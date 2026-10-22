"""DOCX text extraction for the Content Engine import pipeline."""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

from docx import Document as _DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.content.source_document import (
    SourceDocument,
    SourceHeading,
    SourceList,
    SourceParagraph,
    SourceProvenance,
    SourceTable,
    build_source_image,
    document_name,
)


_DRAWING_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_RELATIONSHIP_EMBED = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)


class DocxReader:
    """Extract plain text from DOCX files."""

    def read(self, path: Path) -> str:
        """Read and extract text from all paragraphs of a DOCX file.

        Args:
            path: Path to the DOCX file.

        Returns:
            Extracted text with non-empty paragraphs joined by newlines, or
            an empty string when no text is available.
        """
        document = _DocxDocument(str(path))
        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text
        ]
        return "\n".join(paragraphs)

    def read_document(self, path: Path) -> SourceDocument:
        """Read DOCX body elements in document order without affecting ``read``."""
        document = _DocxDocument(str(path))
        name = document_name(path)
        elements = []
        order = 0
        pending_list: list[str] = []
        pending_list_ordered = False
        pending_list_provenance = None

        def append(element) -> None:
            nonlocal order
            order += 1
            elements.append(element(order))

        def flush_list() -> None:
            nonlocal pending_list, pending_list_provenance
            if not pending_list or pending_list_provenance is None:
                return
            items = tuple(pending_list)
            provenance = pending_list_provenance
            append(
                lambda global_order: SourceList(
                    ordered=pending_list_ordered,
                    items=items,
                    provenance=SourceProvenance(
                        document=provenance.document,
                        order=global_order,
                        position=provenance.position,
                    ),
                )
            )
            pending_list = []
            pending_list_provenance = None

        for position, child in enumerate(document.element.body.iterchildren(), start=1):
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                list_ordered = _paragraph_list_ordered(paragraph)
                image_blobs = tuple(_related_image_blobs(document, child))
                if text and list_ordered is not None:
                    if pending_list and pending_list_ordered != list_ordered:
                        flush_list()
                    if not pending_list:
                        pending_list_ordered = list_ordered
                        pending_list_provenance = SourceProvenance(
                            document=name,
                            order=0,
                            position=position,
                        )
                    pending_list.append(text)
                    # A drawing inside a list item ends this consecutive list
                    # segment so the list remains before its nearby image.
                    if image_blobs:
                        flush_list()
                else:
                    flush_list()
                    if text:
                        level = _heading_level(paragraph)
                        if level is None:
                            append(
                                lambda global_order: SourceParagraph(
                                    text=text,
                                    provenance=SourceProvenance(
                                        document=name,
                                        order=global_order,
                                        position=position,
                                    ),
                                )
                            )
                        else:
                            append(
                                lambda global_order: SourceHeading(
                                    text=text,
                                    level=level,
                                    provenance=SourceProvenance(
                                        document=name,
                                        order=global_order,
                                        position=position,
                                    ),
                                )
                            )
                for blob in image_blobs:
                    append(
                        lambda global_order, blob=blob: build_source_image(
                            document=name,
                            content=blob,
                            provenance=SourceProvenance(
                                document=name,
                                order=global_order,
                                position=position,
                            ),
                        )
                    )
            elif isinstance(child, CT_Tbl):
                flush_list()
                table = Table(child, document)
                rows = tuple(
                    tuple(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                )
                append(
                    lambda global_order: SourceTable(
                        rows=rows,
                        provenance=SourceProvenance(
                            document=name,
                            order=global_order,
                            position=position,
                        ),
                    )
                )
                for blob in _related_image_blobs(document, child):
                    append(
                        lambda global_order, blob=blob: build_source_image(
                            document=name,
                            content=blob,
                            provenance=SourceProvenance(
                                document=name,
                                order=global_order,
                                position=position,
                            ),
                        )
                    )
        flush_list()
        return SourceDocument(name=name, source_format="docx", elements=tuple(elements))


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = str(getattr(paragraph.style, "name", "")).strip().lower()
    if not style_name.startswith("heading "):
        return None
    try:
        level = int(style_name.rsplit(" ", 1)[1])
    except ValueError:
        return None
    return level if level > 0 else None


def _paragraph_list_ordered(paragraph: Paragraph) -> bool | None:
    """Best-effort list detection; numbering without a clear style is unordered."""
    style_name = str(getattr(paragraph.style, "name", "")).strip().lower()
    if style_name.startswith("list number"):
        return True
    if style_name.startswith("list bullet"):
        return False
    if style_name.startswith("list "):
        return False
    paragraph_properties = getattr(paragraph._p, "pPr", None)
    if getattr(paragraph_properties, "numPr", None) is not None:
        return False
    return None


def _related_image_blobs(document, element) -> Iterator[bytes]:
    for descendant in element.iter():
        if descendant.tag != _DRAWING_BLIP:
            continue
        relationship_id = descendant.get(_RELATIONSHIP_EMBED)
        if not relationship_id:
            continue
        related_part = document.part.related_parts.get(relationship_id)
        blob = getattr(related_part, "blob", None)
        if blob:
            yield bytes(blob)
