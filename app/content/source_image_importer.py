"""Extract embedded source images and install them into generated courses."""

from __future__ import annotations

import hashlib
import io
import logging
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

from PIL import Image, UnidentifiedImageError
from docx import Document as _DocxDocument
from pptx import Presentation as _PptxPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader as _PdfReader

from app.content.contract import (
    COURSE_COVER_FILENAMES,
    COURSE_COVER_STEM,
    LESSON_IMAGE_FILENAMES,
    LESSON_IMAGE_STEM,
    LESSON_JSON_FILENAME,
)

_logger = logging.getLogger(__name__)

_DRAWING_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_RELATIONSHIP_EMBED = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
_SUPPORTED_SOURCE_SUFFIXES = frozenset({".docx", ".pptx", ".pdf"})
_MAX_IMAGES = 50
_MAX_IMAGE_BYTES = 10 * 1024 * 1024
_MAX_TOTAL_BYTES = 50 * 1024 * 1024
_MAX_PIXELS = 25_000_000
_MIN_WIDTH = 120
_MIN_HEIGHT = 80
_MIN_AREA = 20_000


@dataclass(frozen=True)
class ExtractedSourceImage:
    """One validated web image extracted from an uploaded source."""

    content: bytes
    extension: str
    width: int
    height: int


@dataclass(frozen=True)
class SourceImageImportResult:
    """Summary of images installed into one generated course."""

    extracted_count: int
    cover_installed: bool
    lesson_images_installed: int


class SourceImageImportService:
    """Extract source images and populate fixed Content Engine image slots."""

    def import_into_course(
        self,
        source_path: Path,
        course_directory: Path,
    ) -> SourceImageImportResult:
        """Install the first image as cover and distribute images over lessons."""
        images = self.extract(source_path)
        if not images:
            return SourceImageImportResult(0, False, 0)

        course_directory = course_directory.resolve()
        if not course_directory.is_dir():
            raise NotADirectoryError(f"Course directory does not exist: {course_directory}")

        _write_image_slot(
            course_directory,
            stem=COURSE_COVER_STEM,
            slot_filenames=COURSE_COVER_FILENAMES,
            image=images[0],
        )

        lesson_directories = sorted(
            path.parent
            for path in course_directory.glob(f"*/{LESSON_JSON_FILENAME}")
            if path.is_file()
        )
        installed = 0
        for lesson_index, image in _distributed_lesson_images(
            images,
            len(lesson_directories),
        ):
            _write_image_slot(
                lesson_directories[lesson_index],
                stem=LESSON_IMAGE_STEM,
                slot_filenames=LESSON_IMAGE_FILENAMES,
                image=image,
            )
            installed += 1

        return SourceImageImportResult(len(images), True, installed)

    def extract(self, source_path: Path) -> tuple[ExtractedSourceImage, ...]:
        """Extract validated, de-duplicated images in source order."""
        suffix = source_path.suffix.lower()
        if suffix not in _SUPPORTED_SOURCE_SUFFIXES:
            return ()

        if suffix == ".docx":
            blobs = _iter_docx_images(source_path)
        elif suffix == ".pptx":
            blobs = _iter_pptx_images(source_path)
        else:
            blobs = _iter_pdf_images(source_path)

        images: list[ExtractedSourceImage] = []
        seen: set[str] = set()
        total_bytes = 0
        for blob in blobs:
            if len(images) >= _MAX_IMAGES:
                break
            image = _normalize_image(blob)
            if image is None:
                continue
            digest = hashlib.sha256(image.content).hexdigest()
            if digest in seen:
                continue
            if total_bytes + len(image.content) > _MAX_TOTAL_BYTES:
                break
            seen.add(digest)
            images.append(image)
            total_bytes += len(image.content)
        return tuple(images)


def _iter_docx_images(path: Path) -> Iterator[bytes]:
    document = _DocxDocument(str(path))
    for element in document.element.body.iter():
        if element.tag != _DRAWING_BLIP:
            continue
        relationship_id = element.get(_RELATIONSHIP_EMBED)
        if not relationship_id:
            continue
        related_part = document.part.related_parts.get(relationship_id)
        blob = getattr(related_part, "blob", None)
        if blob:
            yield bytes(blob)


def _iter_pptx_images(path: Path) -> Iterator[bytes]:
    presentation = _PptxPresentation(str(path))
    for slide in presentation.slides:
        yield from _iter_pptx_shapes(slide.shapes)


def _iter_pptx_shapes(shapes: Iterable[object]) -> Iterator[bytes]:
    for shape in shapes:
        shape_type = getattr(shape, "shape_type", None)
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pptx_shapes(getattr(shape, "shapes", ()))
            continue
        if shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        image = getattr(shape, "image", None)
        blob = getattr(image, "blob", None)
        if blob:
            yield bytes(blob)


def _iter_pdf_images(path: Path) -> Iterator[bytes]:
    reader = _PdfReader(str(path))
    for page in reader.pages:
        try:
            images = page.images
        except Exception:
            _logger.warning("Could not enumerate images on one PDF page", exc_info=True)
            continue
        for image in images:
            data = getattr(image, "data", None)
            if data:
                yield bytes(data)


def _normalize_image(blob: bytes) -> ExtractedSourceImage | None:
    if not blob or len(blob) > _MAX_IMAGE_BYTES:
        return None
    try:
        with Image.open(io.BytesIO(blob)) as opened:
            width, height = opened.size
            if (
                width < _MIN_WIDTH
                or height < _MIN_HEIGHT
                or width * height < _MIN_AREA
                or width * height > _MAX_PIXELS
            ):
                return None
            opened.load()
            image_format = (opened.format or "").upper()
            if image_format in {"JPEG", "PNG", "WEBP"}:
                extension = {
                    "JPEG": ".jpg",
                    "PNG": ".png",
                    "WEBP": ".webp",
                }[image_format]
                content = blob
            else:
                converted = opened.convert("RGBA" if "A" in opened.getbands() else "RGB")
                output = io.BytesIO()
                converted.save(output, format="PNG", optimize=True)
                content = output.getvalue()
                extension = ".png"
    except (OSError, UnidentifiedImageError, ValueError, Image.DecompressionBombError):
        return None

    if not content or len(content) > _MAX_IMAGE_BYTES:
        return None
    return ExtractedSourceImage(content, extension, width, height)


def _distributed_lesson_images(
    images: tuple[ExtractedSourceImage, ...],
    lesson_count: int,
) -> Iterator[tuple[int, ExtractedSourceImage]]:
    if not images or lesson_count <= 0:
        return
    count = min(len(images), lesson_count)
    for image_index in range(count):
        lesson_index = (image_index * lesson_count) // count
        yield lesson_index, images[image_index]


def _write_image_slot(
    directory: Path,
    *,
    stem: str,
    slot_filenames: Iterable[str],
    image: ExtractedSourceImage,
) -> Path:
    target = directory / f"{stem}{image.extension}"
    file_descriptor, temporary_name = tempfile.mkstemp(
        dir=directory,
        prefix=f".{stem}.",
        suffix=".tmp",
    )
    temporary_path = Path(temporary_name)
    try:
        os.fchmod(file_descriptor, 0o640)
        with os.fdopen(file_descriptor, "wb") as handle:
            handle.write(image.content)
            handle.flush()
            os.fsync(handle.fileno())
        for filename in slot_filenames:
            candidate = directory / filename
            if candidate != target and candidate.is_file():
                candidate.unlink()
        os.replace(temporary_path, target)
    except OSError:
        if temporary_path.exists():
            temporary_path.unlink()
        raise
    return target
