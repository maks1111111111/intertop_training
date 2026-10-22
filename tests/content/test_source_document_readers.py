"""Focused coverage for Source Document IR extraction APIs."""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from PIL import Image
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches

from app.content.docx_reader import DocxReader
from app.content.pdf_reader import PdfReader
from app.content.pptx_reader import PptxReader
from app.content.source_document import (
    SourceBoundary,
    SourceHeading,
    SourceImage,
    SourceList,
    SourceParagraph,
    SourceTable,
)


def _png_bytes(color: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (640, 360), color=color).save(output, format="PNG")
    return output.getvalue()


class SourceDocumentReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    @patch("app.content.pdf_reader._PypdfReader")
    def test_pdf_keeps_page_and_image_provenance(
        self, mock_reader_factory: MagicMock
    ) -> None:
        first_page = MagicMock()
        first_page.extract_text.return_value = "Page one"
        first_page.images = [SimpleNamespace(data=_png_bytes("red"))]
        second_page = MagicMock()
        second_page.extract_text.return_value = "Page two"
        second_page.images = []
        reader = MagicMock()
        reader.pages = [first_page, second_page]
        mock_reader_factory.return_value = reader

        document = PdfReader().read_document(Path("/tmp/guide.pdf"))

        self.assertEqual(document.name, "guide.pdf")
        self.assertEqual(document.source_format, "pdf")
        self.assertEqual(
            [type(item) for item in document.elements],
            [SourceBoundary, SourceParagraph, SourceImage, SourceBoundary, SourceParagraph],
        )
        image = document.elements[2]
        assert isinstance(image, SourceImage)
        self.assertEqual(image.provenance.page, 1)
        self.assertEqual(image.provenance.position, 2)
        self.assertTrue(image.asset_id.startswith("source-"))
        self.assertEqual(
            [item.provenance.order for item in document.elements],
            [1, 2, 3, 4, 5],
        )

    def test_docx_preserves_heading_paragraph_table_list_and_image_order(self) -> None:
        image_path = self.root / "guide.png"
        image_path.write_bytes(_png_bytes("blue"))
        source_path = self.root / "guide.docx"
        source = Document()
        source.add_heading("Introduction", level=1)
        source.add_paragraph("Before table")
        table = source.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Cell A"
        table.rows[0].cells[1].text = "Cell B"
        source.add_paragraph("After table")
        source.add_paragraph("First", style="List Bullet")
        source.add_paragraph("Second", style="List Bullet")
        source.add_picture(str(image_path), width=Inches(4))
        source.save(source_path)

        document = DocxReader().read_document(source_path)

        self.assertEqual(document.source_format, "docx")
        self.assertEqual(
            [type(item) for item in document.elements],
            [SourceHeading, SourceParagraph, SourceTable, SourceParagraph, SourceList, SourceImage],
        )
        heading = document.elements[0]
        table_element = document.elements[2]
        source_list = document.elements[4]
        image = document.elements[5]
        assert isinstance(heading, SourceHeading)
        assert isinstance(table_element, SourceTable)
        assert isinstance(source_list, SourceList)
        assert isinstance(image, SourceImage)
        self.assertEqual(heading.level, 1)
        self.assertEqual(table_element.rows, (("Cell A", "Cell B"),))
        self.assertFalse(source_list.ordered)
        self.assertEqual(source_list.items, ("First", "Second"))
        self.assertEqual(image.provenance.document, "guide.docx")
        self.assertGreater(image.provenance.position, source_list.provenance.position)

    def test_pptx_preserves_slide_boundaries_shape_order_pictures_and_tables(self) -> None:
        image_path = self.root / "slide.png"
        image_path.write_bytes(_png_bytes("green"))
        source_path = self.root / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1)
        )
        textbox.text = "Slide text"
        slide.shapes.add_picture(
            str(image_path), PptxInches(1), PptxInches(2)
        )
        table = slide.shapes.add_table(
            1, 2, PptxInches(1), PptxInches(4), PptxInches(5), PptxInches(1)
        ).table
        table.cell(0, 0).text = "Left"
        table.cell(0, 1).text = "Right"
        presentation.save(source_path)

        first = PptxReader().read_document(source_path)
        second = PptxReader().read_document(source_path)

        self.assertEqual(
            [type(item) for item in first.elements],
            [SourceBoundary, SourceParagraph, SourceImage, SourceTable],
        )
        self.assertEqual(
            [item.provenance.order for item in first.elements],
            [item.provenance.order for item in second.elements],
        )
        first_image = first.elements[2]
        first_table = first.elements[3]
        assert isinstance(first_image, SourceImage)
        assert isinstance(first_table, SourceTable)
        self.assertEqual(first_image.provenance.slide, 1)
        self.assertEqual(first_table.rows, (("Left", "Right"),))
        self.assertEqual(first_image.asset_id, second.elements[2].asset_id)


if __name__ == "__main__":
    unittest.main()
