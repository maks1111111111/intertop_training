"""Tests for automatic image extraction from course source files."""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches

from app.content.source_image_importer import SourceImageImportService


def _png_bytes(color: str, size: tuple[int, int] = (640, 360)) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color=color).save(output, format="PNG")
    return output.getvalue()


def _write_course(course_dir: Path, lessons: int = 3) -> None:
    course_dir.mkdir()
    (course_dir / "course.json").write_text(
        json.dumps({"slug": "generated", "title": "Generated"}),
        encoding="utf-8",
    )
    for index in range(1, lessons + 1):
        lesson_dir = course_dir / f"lesson_{index:02d}"
        lesson_dir.mkdir()
        (lesson_dir / "lesson.json").write_text(
            json.dumps({"order": index, "title": f"Lesson {index}"}),
            encoding="utf-8",
        )


class SourceImageImportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.service = SourceImageImportService()

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_extracts_docx_images_and_installs_course_slots(self) -> None:
        first = self.root / "first.png"
        second = self.root / "second.png"
        first.write_bytes(_png_bytes("red"))
        second.write_bytes(_png_bytes("blue"))
        source = self.root / "source.docx"
        document = Document()
        document.add_paragraph("Course")
        document.add_picture(str(first), width=Inches(4))
        document.add_picture(str(second), width=Inches(4))
        document.save(source)

        course_dir = self.root / "course"
        _write_course(course_dir, lessons=4)
        result = self.service.import_into_course(source, course_dir)

        self.assertEqual(result.extracted_count, 2)
        self.assertTrue(result.cover_installed)
        self.assertEqual(result.lesson_images_installed, 2)
        self.assertTrue((course_dir / "cover.png").is_file())
        self.assertTrue((course_dir / "lesson_01" / "image.png").is_file())
        self.assertTrue((course_dir / "lesson_03" / "image.png").is_file())
        self.assertEqual((course_dir / "cover.png").stat().st_mode & 0o777, 0o640)

    def test_extracts_pptx_picture_in_slide_order(self) -> None:
        first = self.root / "first.png"
        second = self.root / "second.png"
        first.write_bytes(_png_bytes("green"))
        second.write_bytes(_png_bytes("yellow"))
        source = self.root / "source.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_picture(
            str(first), PptxInches(1), PptxInches(1)
        )
        presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_picture(
            str(second), PptxInches(1), PptxInches(1)
        )
        presentation.save(source)

        images = self.service.extract(source)

        self.assertEqual(len(images), 2)
        self.assertNotEqual(images[0].content, images[1].content)

    def test_extracts_image_from_pdf(self) -> None:
        source = self.root / "source.pdf"
        Image.open(io.BytesIO(_png_bytes("purple"))).save(source, format="PDF")

        images = self.service.extract(source)

        self.assertEqual(len(images), 1)
        self.assertGreaterEqual(images[0].width, 120)

    def test_deduplicates_and_ignores_small_images(self) -> None:
        large = self.root / "large.png"
        small = self.root / "small.png"
        large.write_bytes(_png_bytes("black"))
        small.write_bytes(_png_bytes("white", (40, 40)))
        source = self.root / "source.docx"
        document = Document()
        document.add_picture(str(large))
        document.add_picture(str(large))
        document.add_picture(str(small))
        document.save(source)

        images = self.service.extract(source)

        self.assertEqual(len(images), 1)

    def test_unsupported_source_has_no_images(self) -> None:
        source = self.root / "source.txt"
        source.write_text("text", encoding="utf-8")

        result = self.service.extract(source)

        self.assertEqual(result, ())


if __name__ == "__main__":
    unittest.main()
